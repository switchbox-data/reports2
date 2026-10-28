"""Excel workbooks as sources: no extraction (agents read the workbook itself), a stub ``.md`` next to it.

A workbook under the repo's large-file limit is copied next to the other extracts. Above it, the card asks:
commit it anyway (the commit must skip the ``check-added-large-files`` hook) or keep it only in Zotero and
fetch it with ``just fetch-source``. Either way the stub holds the front matter a workbook cannot carry.
"""

from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from lib.sources import frontmatter

LIMIT_BYTES = 600 * 1024  # .pre-commit-config.yaml: check-added-large-files --maxkb=600
NOT_IN_REPO = "not in repo: fetch from Zotero"
SKIP_HINT = "SKIP=check-added-large-files git commit ..."


def too_large(path: Path) -> bool:
    """Whether the workbook exceeds the repo's 600 KB commit limit."""
    return Path(path).stat().st_size > LIMIT_BYTES


def needs_decision(path: Path, workbook_in_repo: bool | None) -> bool:
    """Whether the card must ask about committing: the file is too large and nothing was decided yet."""
    return too_large(path) and workbook_in_repo is None


def copy_into_repo(src: Path, dest: Path) -> bool:
    """Copy the workbook to ``dest`` unless an identical copy is already there; True when it wrote."""
    dest = Path(dest)
    if dest.exists() and dest.read_bytes() == Path(src).read_bytes():
        return False
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dest)
    return True


def sheet_summary(path: Path, max_columns: int = 12) -> str:
    """A markdown list of the workbook's sheets: name, size, and the first non-empty row (usually the headers).

    Empty when openpyxl cannot read the file (an ``.xls`` or a damaged workbook).
    """
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, KeyError, OSError):
        return ""
    lines = []
    try:
        for ws in wb.worksheets:
            first = next((row for row in ws.iter_rows(values_only=True) if any(v is not None for v in row)), ())
            cells = [str(v) for v in first if v is not None]
            shown = ", ".join(cells[:max_columns]) + (", ..." if len(cells) > max_columns else "")
            line = f"- **{ws.title}** ({ws.max_row or 0} rows x {ws.max_column or 0} columns)"
            lines.append(f"{line}: {shown}" if shown else line)
    finally:
        wb.close()
    return "\n".join(lines)


def write_stub(
    path: Path, fields: dict[str, Any], title: str, description: str, workbook_path: Path | None = None
) -> None:
    """Write (or refresh the front matter of) the stub next to a workbook; keeps an existing body.

    With ``workbook_path``, a new stub lists the sheets so a reader knows what the workbook holds before opening it.
    """
    path = Path(path)
    if path.exists():
        frontmatter.update(path, fields)
        return
    body = f"# {title}\n\n{description.strip()}\n"
    sheets = sheet_summary(workbook_path) if workbook_path is not None else ""
    if sheets:
        body += f"\n## Sheets\n\n{sheets}\n"
    frontmatter.write(path, fields, body)


def repo_file_field(workbook_in_repo: bool, filename: str) -> str:
    return filename if workbook_in_repo else NOT_IN_REPO
